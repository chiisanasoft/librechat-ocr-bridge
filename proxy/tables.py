"""HTML tables from OCR output -> Markdown tables and .xlsx workbooks."""

import os
import re
from dataclasses import dataclass, field
from html import unescape
from html.parser import HTMLParser

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

import headers

# Add inferred column headers (日付, 姓, 郵便番号, ...) to tables that have none
AUTO_HEADER = os.environ.get("TABLE_AUTO_HEADER", "true").lower() in ("1", "true", "yes")

TABLE_RE = re.compile(r"<table\b.*?</table>", re.IGNORECASE | re.DOTALL)


@dataclass
class Cell:
    text: str
    colspan: int = 1
    rowspan: int = 1
    header: bool = False


@dataclass
class Table:
    rows: list[list[Cell]] = field(default_factory=list)

    def grid(self) -> tuple[list[list[str]], list[tuple[int, int, int, int]]]:
        """Expands spans into a rectangular grid; returns (grid, merges as r1,c1,r2,c2)."""
        placed: dict[tuple[int, int], str] = {}
        merges = []
        for r, row in enumerate(self.rows):
            c = 0
            for cell in row:
                while (r, c) in placed:
                    c += 1
                for dr in range(cell.rowspan):
                    for dc in range(cell.colspan):
                        placed[(r + dr, c + dc)] = cell.text if dr == dc == 0 else ""
                if cell.rowspan > 1 or cell.colspan > 1:
                    merges.append((r, c, r + cell.rowspan - 1, c + cell.colspan - 1))
                c += cell.colspan
        if not placed:
            return [], []
        n_rows = max(r for r, _ in placed) + 1
        n_cols = max(c for _, c in placed) + 1
        grid = [[placed.get((r, c), "") for c in range(n_cols)] for r in range(n_rows)]
        return grid, merges


class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.table = Table()
        self.cell: Cell | None = None
        self.buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "tr":
            self.table.rows.append([])
        elif tag in ("td", "th"):
            if not self.table.rows:
                self.table.rows.append([])
            self.cell = Cell("", _span(a.get("colspan")), _span(a.get("rowspan")), tag == "th")
            self.buf = []
        elif tag == "br" and self.cell is not None:
            self.buf.append("\n")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            self.cell.text = re.sub(r"[ \t]+", " ", "".join(self.buf)).strip()
            self.table.rows[-1].append(self.cell)
            self.cell = None

    def handle_data(self, data):
        if self.cell is not None:
            self.buf.append(data)


def _span(v) -> int:
    try:
        return max(1, min(int(v), 100))
    except (TypeError, ValueError):
        return 1


def parse_table(html: str) -> Table:
    p = _TableParser()
    p.feed(html)
    p.close()
    return p.table


def add_inferred_header(table: Table) -> bool:
    grid, _ = table.grid()
    labels = headers.infer_headers(grid)
    if not labels:
        return False
    # OCR models tag the first data row as <th>; it is data once a real header exists
    for row in table.rows:
        for cell in row:
            cell.header = False
    table.rows.insert(0, [Cell(label, header=True) for label in labels])
    return True


def to_markdown(table: Table) -> str:
    grid, _ = table.grid()
    if not grid:
        return ""

    def esc(s: str) -> str:
        return unescape(s).replace("|", "\\|").replace("\n", "<br>") or " "

    lines = ["| " + " | ".join(esc(c) for c in grid[0]) + " |",
             "|" + "---|" * len(grid[0])]
    lines += ["| " + " | ".join(esc(c) for c in row) + " |" for row in grid[1:]]
    return "\n".join(lines)


def convert_html_tables(text: str) -> tuple[str, list[Table]]:
    """Replaces each <table> in OCR output with a Markdown table."""
    tables: list[Table] = []

    def repl(m: re.Match) -> str:
        t = parse_table(m.group(0))
        if not t.rows:
            return m.group(0)
        if AUTO_HEADER:
            add_inferred_header(t)
        tables.append(t)
        return "\n\n" + to_markdown(t) + "\n\n"

    return TABLE_RE.sub(repl, text).strip(), tables


def write_xlsx(path: str, sheets: list[tuple[str, Table]]) -> None:
    """One worksheet per (name, table); spans are kept as merged cells."""
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for name, table in sheets:
        ws = wb.create_sheet(_sheet_title(name, wb.sheetnames))
        grid, merges = table.grid()
        header_rows = {r for r, row in enumerate(table.rows) if row and all(c.header for c in row)}
        widths: dict[int, int] = {}
        for r, row in enumerate(grid, 1):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=_typed(unescape(value)))
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if r - 1 in header_rows:
                    cell.font = Font(bold=True)
                longest = max((_display_width(s) for s in value.split("\n")), default=0)
                widths[c] = max(widths.get(c, 0), longest)
        for r1, c1, r2, c2 in merges:
            ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1, end_row=r2 + 1, end_column=c2 + 1)
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 6), 60)
    wb.save(path)


def _typed(value: str):
    """Keeps codes like 03-1234-5678 or 0120 as text; converts plain numbers."""
    if re.fullmatch(r"-?[1-9]\d{0,14}", value) or re.fullmatch(r"-?\d+\.\d+", value):
        return float(value) if "." in value else int(value)
    return value


def _display_width(s: str) -> int:
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)


def _sheet_title(name: str, existing: list[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:28] or "Sheet"
    title, n = base, 2
    while title in existing:
        title, n = f"{base}_{n}", n + 1
    return title
