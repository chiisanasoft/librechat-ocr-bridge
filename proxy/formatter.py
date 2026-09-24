"""OCR -> one cleaned table: merge pages, reshape to a column template with an LLM,
normalize values by rule, and flag every cell the reader should double-check.

Date columns are handled by rule (a small LLM splits "12/5K12/8T" inconsistently):
the first date goes to the template's date column, later dates and code letters
go to the notes column. The LLM reshapes the remaining columns and is only
trusted to move, split and lightly correct values. Its output is
checked against the source row: values that do not occur in the source are
marked as corrections (with the original kept), letters/digits that disappeared
are reported as possible omissions, and format rules flag invalid dates, postal
codes and phone numbers.
"""

import json
import os
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import tables

FORMAT_MODEL_NAME = os.environ.get("FORMAT_MODEL_NAME", "ocr-format")
FORMAT_LLM = os.environ.get("FORMAT_LLM", "nemotron-nano-9b-ja:latest")
FORMAT_TEMPLATE = [c.strip() for c in os.environ.get(
    "FORMAT_TEMPLATE", "日付,姓,名,郵便番号,住所,電話番号,備考").split(",") if c.strip()]
FORMAT_BATCH_ROWS = int(os.environ.get("FORMAT_BATCH_ROWS", "15"))
NOTES_COLUMN = "備考"

SYSTEM_PROMPT = """あなたはOCRで読み取った表を、指定された列の表に整形する担当です。次のルールを必ず守ってください。
- 入力の各行につき、必ず1行を出力する。src には入力行の番号を入れる。
- 値は入力行の文字列から取る。入力にない情報を推測して作らない。該当する値がなければ空文字にする。
- 直してよいのは、形の似た文字の誤認識など明らかなOCRの誤りだけ。
- 1つのセルに複数の値（例: 日付が2つ）が入っていれば分割する。出力の列に入らない値は「備考」に入れる。
- 日付に付いた記号やアルファベット（例: 12/5K の K）、チェック記号、メモなどは捨てずに「備考」に入れる。
- 氏名が1つのセルにあり、出力に「姓」「名」の列があれば分割する。
- 郵便番号・電話番号・住所は、入力の見出し推定に関係なく、値の内容から正しい列に入れる。"""


# ---------------------------------------------------------------- template

def parse_template(text: str) -> list[str]:
    """Column template from the chat ("列: 日付, 氏名, 住所"), else the configured default."""
    m = re.search(r"(?:列|項目|columns?)\s*[:：]\s*([^\n]+)", text, re.IGNORECASE)
    if m:
        cols = [c.strip() for c in re.split(r"[,、，/／]", m.group(1)) if c.strip()]
        if cols:
            return cols
    return list(FORMAT_TEMPLATE)


def column_kind(name: str) -> str | None:
    if re.search(r"日付|年月日|日時|date", name, re.IGNORECASE):
        return "date"
    if re.search(r"郵便|〒|zip|postal", name, re.IGNORECASE):
        return "postal"
    if re.search(r"電話|tel|phone|携帯", name, re.IGNORECASE):
        return "phone"
    return None


# ---------------------------------------------------------------- merging

@dataclass
class SourceRow:
    page: int
    cells: dict[str, str]  # source header -> value

    def text(self) -> str:
        return " ".join(v for v in self.cells.values() if v)


def merge_pages(pages: list[list[tables.Table]]) -> tuple[list[str], list[SourceRow]]:
    """Joins the tables of all pages into one list of rows keyed by (inferred) header."""
    columns: list[str] = []
    rows: list[SourceRow] = []
    for page_no, page_tables in enumerate(pages, 1):
        for t in page_tables:
            grid, _ = t.grid()
            if not grid:
                continue
            has_header = bool(t.rows) and all(c.header for c in t.rows[0])
            header = grid[0] if has_header else [f"列{i + 1}" for i in range(len(grid[0]))]
            header = [h or f"列{i + 1}" for i, h in enumerate(header)]
            for h in header:
                if h not in columns:
                    columns.append(h)
            for values in grid[1:] if has_header else grid:
                if not any(v.strip() for v in values) or values == header:
                    continue  # blank rows and header rows repeated on later pages
                rows.append(SourceRow(page_no, dict(zip(header, values))))
    return columns, rows


# ---------------------------------------------------------------- dates by rule

DATE_TOKEN = re.compile(r"(\d{1,2})\s*[/・.．]\s*(\d{1,2})(?!\d)\s*([A-Za-z]?)")


def is_date_source(header: str) -> bool:
    return column_kind(header) == "date"


def split_dates(row: SourceRow) -> tuple[str, list[str]]:
    """First date of the row's date columns, plus notes for suffixes, later dates and leftovers."""
    first, notes = "", []
    for header, value in row.cells.items():
        if not is_date_source(header) or not value.strip():
            continue
        v = unicodedata.normalize("NFKC", value)
        for m in DATE_TOKEN.finditer(v):
            date, code = f"{int(m.group(1))}/{int(m.group(2))}", m.group(3)
            if not first:
                first = date
                if code:
                    notes.append(code)
            else:
                notes.append(date + code)
        leftover = DATE_TOKEN.sub(" ", v).strip(" ・,、")
        if leftover:
            notes.append(leftover)
    return first, notes


# ---------------------------------------------------------------- LLM reshaping

def row_schema(template: list[str]) -> dict:
    keys = template + ([NOTES_COLUMN] if NOTES_COLUMN not in template else [])
    item = {"type": "object",
            "properties": {"src": {"type": "integer"}, **{k: {"type": "string"} for k in keys}},
            "required": ["src", *keys]}
    return {"type": "object", "properties": {"rows": {"type": "array", "items": item}},
            "required": ["rows"]}


def rule_date_columns(template: list[str], source_columns: list[str]) -> bool:
    return any(column_kind(c) == "date" for c in template) and any(map(is_date_source, source_columns))


async def reshape_batch(client, upstream: str, headers: dict, template: list[str],
                        source_columns: list[str], batch: list[tuple[int, SourceRow]]) -> dict[int, dict]:
    if rule_date_columns(template, source_columns):
        # Dates are filled by rule in finalize(); keep them away from the LLM
        template = [c for c in template if column_kind(c) != "date"]
        source_columns = [c for c in source_columns if not is_date_source(c)]
    lines = [f"{i}: " + json.dumps([row.cells.get(c, "") for c in source_columns], ensure_ascii=False)
             for i, row in batch]
    user = (f"出力する列: {', '.join(template)}\n"
            f"入力の列（見出しは推定）: {json.dumps(source_columns, ensure_ascii=False)}\n"
            f"入力行:\n" + "\n".join(lines))
    req = {"model": FORMAT_LLM, "stream": False, "think": False, "format": row_schema(template),
           "options": {"temperature": 0, "num_ctx": 8192},
           "messages": [{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": user}]}
    r = await client.post(f"{upstream}/api/chat", json=req, headers=headers)
    r.raise_for_status()
    out = json.loads(r.json()["message"]["content"])
    wanted = {i for i, _ in batch}
    return {row["src"]: row for row in out.get("rows", []) if row.get("src") in wanted}


# ---------------------------------------------------------------- rules & checks

def canon(s: str) -> str:
    s = unicodedata.normalize("NFKC", s).lower()
    return re.sub(r"[\s・/.\-ー－()（）,、]", "", s)


def normalize(value: str, kind: str | None) -> str:
    v = unicodedata.normalize("NFKC", value).strip()
    if kind == "date":
        m = re.fullmatch(r"(\d{1,4})\s*[/・.年-]\s*(\d{1,2})(?:\s*[/・.月-]\s*(\d{1,2}))?\s*日?", v)
        if m:
            v = "/".join(g for g in m.groups() if g)
    elif kind == "postal":
        digits = re.sub(r"\D", "", v)
        if len(digits) == 7:
            v = f"{digits[:3]}-{digits[3:]}"
    elif kind == "phone":
        v = re.sub(r"\s+", "", v)
    return v


def invalid(value: str, kind: str | None) -> str | None:
    if not value or kind is None:
        return None
    if kind == "date":
        parts = [int(p) for p in value.split("/")] if re.fullmatch(r"\d{1,4}(/\d{1,2}){1,2}", value) else []
        month, day = (parts[-2], parts[-1]) if len(parts) >= 2 else (0, 0)
        if not (1 <= month <= 12 and 1 <= day <= 31):
            return "日付の形式ではありません"
    elif kind == "postal" and not re.fullmatch(r"\d{3}-\d{4}", value):
        return "郵便番号の形式ではありません"
    elif kind == "phone" and not re.fullmatch(r"0\d{1,4}-\d{1,4}-\d{3,4}|0\d{9,10}", value):
        return "電話番号の形式ではありません"
    return None


@dataclass
class Issue:
    row: int  # 1-based output row
    column: str
    kind: str  # 補正 / 要確認 / 欠落の可能性 / 未整形
    detail: str


@dataclass
class Result:
    columns: list[str]
    rows: list[list[str]]
    issues: list[Issue] = field(default_factory=list)

    def cell_issues(self) -> dict[tuple[int, str], Issue]:
        return {(i.row, i.column): i for i in self.issues}


def fallback_row(row: SourceRow, template: list[str]) -> dict:
    """Maps source columns whose (inferred) header matches a template column."""
    return {c: row.cells.get(c, "") for c in template}


def finalize(template: list[str], source: list[SourceRow], shaped: dict[int, dict],
             source_columns: list[str] | None = None) -> Result:
    columns = list(template)
    date_col = next((c for c in template if column_kind(c) == "date"), None)
    by_rule = rule_date_columns(template, source_columns or [])
    if by_rule:
        for i, src in enumerate(source):
            if i in shaped:
                first, extra = split_dates(src)
                shaped[i][date_col] = first
                shaped[i][NOTES_COLUMN] = ", ".join(
                    p for p in [*extra, shaped[i].get(NOTES_COLUMN, "")] if p)
    notes_used = NOTES_COLUMN not in template and any(
        (shaped.get(i) or {}).get(NOTES_COLUMN) for i in range(len(source)))
    if notes_used:
        columns.append(NOTES_COLUMN)
    kinds = {c: column_kind(c) for c in columns}
    result = Result(columns, [])

    for i, src in enumerate(source):
        n = i + 1
        out = shaped.get(i)
        if out is None:
            out = fallback_row(src, columns)
            result.issues.append(Issue(n, columns[0], "未整形", "LLMの出力がなかったため、元の列をそのまま使いました"))
        src_text = src.text()
        src_canon = canon(src_text)
        values = []
        for c in columns:
            v = normalize(str(out.get(c) or ""), kinds[c])
            if v and canon(v) and canon(v) not in src_canon:
                result.issues.append(Issue(n, c, "補正", f"元の行: {src_text}"))
            problem = invalid(v, kinds[c])
            if problem:
                result.issues.append(Issue(n, c, "要確認", problem))
            values.append(v)
        missing = Counter(ch for ch in canon(src_text) if ch.isascii() and ch.isalnum())
        missing -= Counter(ch for ch in canon(" ".join(values)) if ch.isascii() and ch.isalnum())
        if missing:
            result.issues.append(Issue(n, columns[-1], "欠落の可能性",
                                       f"元の行にあった「{''.join(sorted(missing.elements()))}」が見当たりません（元の行: {src_text}）"))
        result.rows.append(values)
    return result


# ---------------------------------------------------------------- output

MARKS = {"補正": "✎", "要確認": "⚠", "欠落の可能性": "⚠", "未整形": "⚠"}


def to_markdown(result: Result) -> str:
    issues = result.cell_issues()

    def cell(n: int, c: str, v: str) -> str:
        v = v.replace("|", "\\|").replace("\n", "<br>")
        issue = issues.get((n, c))
        return (f"{v} {MARKS[issue.kind]}" if v else MARKS[issue.kind]) if issue else (v or " ")

    lines = ["| # | " + " | ".join(result.columns) + " |", "|---|" + "---|" * len(result.columns)]
    for n, row in enumerate(result.rows, 1):
        lines.append(f"| {n} | " + " | ".join(cell(n, c, v) for c, v in zip(result.columns, row)) + " |")
    return "\n".join(lines)


def issues_markdown(result: Result, limit: int = 40) -> str:
    if not result.issues:
        return "確認が必要な箇所はありません。"
    lines = [f"- {i.row}行目「{i.column}」{MARKS[i.kind]} {i.kind}: {i.detail}" for i in result.issues[:limit]]
    if len(result.issues) > limit:
        lines.append(f"- ほか {len(result.issues) - limit} 件（Excel の「確認リスト」シートを参照）")
    return "\n".join(lines)


FILLS = {"補正": PatternFill("solid", fgColor="FFF2B3"), "要確認": PatternFill("solid", fgColor="FFD6D6"),
         "欠落の可能性": PatternFill("solid", fgColor="FFD6D6"), "未整形": PatternFill("solid", fgColor="FFD6D6")}


def write_workbook(path: str, result: Result, originals: list[tuple[str, tables.Table]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "整形済み"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    kinds = {c: column_kind(c) for c in result.columns}
    issues = result.cell_issues()
    for c, name in enumerate(result.columns, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.border = Font(bold=True), border
    for n, row in enumerate(result.rows, 1):
        for c, (name, value) in enumerate(zip(result.columns, row), 1):
            # Codes stay text so leading zeros survive
            cell = ws.cell(row=n + 1, column=c, value=value if kinds[name] else tables._typed(value))
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            issue = issues.get((n, name))
            if issue:
                cell.fill = FILLS[issue.kind]
                cell.comment = Comment(f"{issue.kind}: {issue.detail}", "OCR Bridge")
    for c, name in enumerate(result.columns, 1):
        width = max([tables._display_width(name)] + [tables._display_width(r[c - 1]) for r in result.rows])
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 6), 50)
    ws.freeze_panes = "A2"

    log = wb.create_sheet("確認リスト")
    log.append(["行", "列", "種類", "内容"])
    for cell in log[1]:
        cell.font = Font(bold=True)
    for i in result.issues:
        log.append([i.row, i.column, i.kind, i.detail])
    for col, width in zip("ABCD", (6, 12, 12, 80)):
        log.column_dimensions[col].width = width

    for name, table in originals:
        tables.add_table_sheet(wb, f"OCR原本_{name}", table)
    wb.save(path)
